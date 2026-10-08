import tempfile
import unittest
from pathlib import Path
from extract import extract, split_text


class ExtractTests(unittest.TestCase):
    def test_no_text_loss(self):
        text = "pengetahuan organisasi " * 400
        chunks = split_text(text, "Teks")
        self.assertGreater(len(chunks), 1)
        self.assertEqual("".join(c["content"].replace(" ", "") for c in chunks), text.replace(" ", ""))
        self.assertTrue(all(len(c["content"]) <= 1400 for c in chunks))

    def test_docx_and_utf8(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "sumber.docx"
            doc = Document()
            doc.add_paragraph("Kebijakan akses dokumen internal.")
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "Persetujuan administrator"
            doc.save(path)
            result = extract(path, ".docx")
            self.assertEqual(len(result["chunks"]), 2)
            path = Path(folder) / "sumber.md"
            path.write_text("Informasi résumé berbahasa Indonesia", encoding="utf-8")
            self.assertIn("résumé", extract(path, ".md")["chunks"][0]["content"])

    def test_empty_invalid_and_scan(self):
        from pypdf import PdfWriter
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "kosong.txt"
            path.write_text(" ")
            with self.assertRaises(ValueError):
                extract(path, ".txt")
            path.write_bytes(b"\xff")
            with self.assertRaises(UnicodeError):
                extract(path, ".txt")
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            pdf = Path(folder) / "scan.pdf"
            writer.write(pdf)
            with self.assertRaisesRegex(ValueError, "OCR"):
                extract(pdf, ".pdf")


if __name__ == "__main__":
    unittest.main()
