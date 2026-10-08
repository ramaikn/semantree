"""Ekstraksi lokal. Tidak mengikuti URL atau menjalankan isi dokumen."""
import json
import re
import sys
import zipfile
from pathlib import Path

MAX_TEXT = 1_500_000
MAX_EXPANDED = 48_000_000


def split_text(text: str, location: str) -> list[dict]:
    text = text.replace("\x00", "").replace("\r\n", "\n").strip()
    if not text:
        return []
    result = []
    start = 0
    while start < len(text):
        end = min(start + 1400, len(text))
        if end < len(text):
            boundary = max(text.rfind("\n", start + 700, end), text.rfind(" ", start + 700, end))
            if boundary > start:
                end = boundary
        result.append({"content": text[start:end].strip(), "location": f"{location}, karakter {start + 1}–{end}"})
        start = end
    return result


def extract(path: Path, suffix: str) -> dict:
    suffix = suffix.lower()
    chunks = []
    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8-sig", errors="strict")
        # Teks kontrol selain tab/newline tidak disajikan ke browser/model.
        if re.search(r"[\x01-\x08\x0b\x0c\x0e-\x1f]", text):
            raise ValueError("Dokumen mengandung karakter kontrol. Gunakan teks UTF-8.")
        chunks = split_text(text, "Teks")
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("PDF terenkripsi tidak didukung.")
        if len(reader.pages) > 300:
            raise ValueError("Batas PDF adalah 300 halaman.")
        for number, page in enumerate(reader.pages, 1):
            chunks.extend(split_text(page.extract_text() or "", f"Halaman {number}"))
            if sum(len(c["content"]) for c in chunks) > MAX_TEXT:
                raise ValueError("Teks melebihi batas 1,5 juta karakter.")
    elif suffix == ".docx":
        from docx import Document
        with zipfile.ZipFile(path) as archive:
            info = archive.infolist()
            if len(info) > 3000 or sum(i.file_size for i in info) > MAX_EXPANDED:
                raise ValueError("Ukuran isi DOCX melebihi batas aman.")
        doc = Document(path)
        for number, paragraph in enumerate(doc.paragraphs, 1):
            chunks.extend(split_text(paragraph.text, f"Paragraf {number}"))
        for number, table in enumerate(doc.tables, 1):
            text = "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            chunks.extend(split_text(text, f"Tabel {number}"))
    else:
        raise ValueError("Format belum didukung. Gunakan PDF, DOCX, TXT, atau Markdown.")
    length = sum(len(chunk["content"]) for chunk in chunks)
    if not length:
        raise ValueError("Tidak ada teks yang dapat diekstrak. PDF hasil scan memerlukan OCR dan belum didukung.")
    if length > MAX_TEXT or len(chunks) > 3000:
        raise ValueError("Teks atau jumlah potongan melebihi batas aman.")
    return {"chunks": chunks, "textLength": length}


if __name__ == "__main__":
    try:
        result = extract(Path(sys.argv[1]), sys.argv[2])
        sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False).encode("utf-8"))
    except Exception as error:
        message = str(error) if isinstance(error, (ValueError, UnicodeError)) else "Dokumen rusak atau tidak dapat diekstrak."
        sys.stdout.buffer.write(json.dumps({"error": message}, ensure_ascii=False).encode("utf-8"))
        sys.exit(1)
